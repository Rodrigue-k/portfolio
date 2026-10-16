"""Update the French editable CV and regenerate both published PDF versions."""
from pathlib import Path
from html import escape
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph as DocxParagraph
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader
import pypdfium2, runpy

ROOT=Path(__file__).resolve().parents[3]
source=ROOT/'docs/cv/source/CV_Rodrigue_Koudakpo_Developpeur.docx'
doc=Document(source)
def after(anchor,text,style=None):
    element=OxmlElement('w:p');anchor._p.addnext(element)
    p=DocxParagraph(element,anchor._parent);p.style=style or anchor.style;p.add_run(text);return p

for p in doc.paragraphs:
    if p.text.startswith("DÉVELOPPEUR D'APPLICATIONS"):
        p.runs[0].text='DÉVELOPPEUR FLUTTER & FULL STACK · FORMATEUR'
        for r in p.runs[1:]:r.text=''
    if p.text.startswith('Développeur autodidacte :'):
        p.text='Développeur autodidacte depuis 2019, spécialisé Flutter depuis 2022. Je conçois et livre des applications mobiles et web de l’idée au déploiement : connectivité variable, paiements mobiles, utilisateurs non techniques. Associé chez DTC et développeur indépendant, je suis également formateur Python à domicile et en ligne, notamment pour des besoins professionnels de gestion de données.'

if not any(p.text.startswith('Formateur indépendant') for p in doc.paragraphs):
    heading=next(p for p in doc.paragraphs if p.text=='EXPÉRIENCE PROFESSIONNELLE')
    role_style=next(p.style for p in doc.paragraphs if p.text.startswith('Associé &'))
    company_style=next(p.style for p in doc.paragraphs if p.text.startswith('Darollo Technologies'))
    bullet_style=next(p.style for p in doc.paragraphs if p.text.startswith('Développement et maintenance'))
    a=after(heading,'Formateur indépendant\tDepuis 2025',role_style)
    a=after(a,'Cours particuliers à domicile et en ligne',company_style)
    a=after(a,'Formation Python adaptée aux objectifs des apprenants, notamment professionnels : programmation et gestion de données.',bullet_style)
    after(a,'Conception de supports et d’exercices pratiques ; préparation d’une formation de groupe en club universitaire.',bullet_style)
if not any(p.text.startswith('KOD8 —') for p in doc.paragraphs):
    heading=next(p for p in doc.paragraphs if p.text=='PROJETS SÉLECTIONNÉS')
    bullet_style=next(p.style for p in doc.paragraphs if p.text.startswith('ciForms'))
    after(heading,'KOD8 — plateforme de formation conçue et développée de bout en bout : parcours Python de 14 modules, initiation à l’IA et correction automatique des exercices. kod8.rodriguekoudakpo.com',bullet_style)
doc.save(source)

for name,file in [('Calibri','calibri.ttf'),('Calibri-Bold','calibrib.ttf'),('Calibri-Italic','calibrii.ttf')]:pdfmetrics.registerFont(TTFont(name,str(Path('C:/Windows/Fonts')/file)))
pdfmetrics.registerFontFamily('Calibri',normal='Calibri',bold='Calibri-Bold',italic='Calibri-Italic')
body=ParagraphStyle('body',fontName='Calibri',fontSize=9,leading=10.3,spaceAfter=1)
heading=ParagraphStyle('heading',parent=body,fontName='Calibri-Bold',fontSize=11,leading=13,spaceBefore=6,spaceAfter=3,keepWithNext=True)
role=ParagraphStyle('role',parent=body,fontName='Calibri-Bold',fontSize=10,leading=11.5,spaceBefore=3,keepWithNext=True)
name=ParagraphStyle('name',parent=body,fontName='Calibri-Bold',fontSize=27,leading=30,spaceAfter=2)
subtitle=ParagraphStyle('subtitle',parent=body,fontName='Calibri-Bold',fontSize=11,leading=14,spaceAfter=3)
sections={'PROFIL','COMPÉTENCES TECHNIQUES','EXPÉRIENCE PROFESSIONNELLE','PROJETS SÉLECTIONNÉS','FORMATION & CERTIFICATIONS'}
story=[]
for i,p in enumerate(doc.paragraphs):
    t=p.text.strip()
    if not t:continue
    style=name if i==0 else subtitle if i==1 else heading if t in sections else role if '\t' in t and any(t.startswith(a) for a in ['Formateur','Associé','Développeur','Commercial','Opérateur']) else body
    text=escape(t).replace('\t',' · ')
    if 'kod8.rodriguekoudakpo.com' in text:text=text.replace('kod8.rodriguekoudakpo.com','<link href="https://kod8.rodriguekoudakpo.com/">kod8.rodriguekoudakpo.com</link>')
    story.append(Paragraph(text,style))
out=ROOT/'public/cv/CV_Rodrigue_Koudakpo_Developpeur.pdf'
SimpleDocTemplate(str(out),pagesize=A4,leftMargin=40,rightMargin=40,topMargin=24,bottomMargin=24,title='Komi Rodrigue Koudakpo — Développeur Flutter & Full Stack · Formateur',author='Komi Rodrigue Koudakpo').build(story)
qa=ROOT/'tmp/cv-qa';qa.mkdir(parents=True,exist_ok=True)
for lang,path in [('fr',out)]:
    pdf=pypdfium2.PdfDocument(str(path))
    for i in range(len(pdf)):pdf[i].render(scale=1.6).to_pil().save(str(qa/f'{lang}-{i+1}.png'))
    print(lang,'pages:',len(pdf))
runpy.run_path(str(ROOT/'docs/cv/source/build_english_cv.py'))
